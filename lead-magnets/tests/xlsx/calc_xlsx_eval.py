"""Evaluates public/ferramentas/arquivos/calculadora-cac.xlsx for a list of input scenarios.

    echo '[{...scenario...}]' | python3 tests/calc_xlsx_eval.py

Scenario keys match calculator/calc.js inputs (percentages as 0–100). Prints a JSON list of
outputs, with "–" (text) mapped to null, so tests can compare against the web calculator.
"""
import json
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent / "xlsx"))
from paths import DEMO_JSON, OUT  # noqa: E402
from xlsx_cache import ERRORS, recalculate_copy  # noqa: E402

SRC = OUT / "calculadora-cac.xlsx"
SCALARS = {"invest": ("in_invest", 1), "fixos": ("in_fixos", 1), "salario": ("in_salario", 1), "pctSec": ("in_pct_sec", 100),
           "leads": ("in_leads", 1), "qual": ("in_qual", 1), "agend": ("in_agend", 1), "comp": ("in_comp", 1),
           "fech": ("in_fech", 1), "volta": ("in_volta", 100), "retorno": ("in_retorno", 1), "anos": ("in_anos", 1)}
OUTS = ["cpl", "custo_qual", "custo_agend", "custo_comp", "cac_ads", "cac_real", "conv1", "conv2", "conv3", "conv4",
        "conv_total", "um_em_cada", "receita", "lucro", "custo_total", "roi", "roas", "ltv", "ltv_cac",
        "onde1", "onde2", "onde3", "onde4", "mais_anuncios", "maior"]
PROC_FIRST, PROC_LAST = 24, 33


def cell(wb, name):
    sheet, ref = next(iter(wb.defined_names[name].destinations))
    return wb[sheet][ref.replace("$", "")]


def main():
    scenarios = json.load(sys.stdin)
    results = []
    tmp = Path(tempfile.mkdtemp(prefix="calc-eval-"))
    try:
        for i, s in enumerate(scenarios):
            wb = load_workbook(SRC)
            for key, (n, scale) in SCALARS.items():
                v = s.get(key)
                cell(wb, n).value = None if v in (None, "") else (v / scale if scale != 1 else v)
            if not s.get("anos"):
                cell(wb, "in_anos").value = 1
            ws = wb["Calculadora"]
            procs = s.get("procs") or []
            for k, r in enumerate(range(PROC_FIRST, PROC_LAST + 1)):
                p = procs[k] if k < len(procs) else {}
                ws[f"A{r}"] = p.get("nome") or None
                ws[f"B{r}"] = p.get("fech") if p.get("fech") not in (None, "") else None
                ws[f"C{r}"] = p.get("ticket") if p.get("ticket") not in (None, "") else None
                ws[f"D{r}"] = p["margem"] / 100 if p.get("margem") not in (None, "") else None
            path = tmp / f"s{i}.xlsx"
            wb.save(path)
            (tmp / f"w{i}").mkdir()
            calc = load_workbook(recalculate_copy(path, tmp / f"w{i}"), data_only=True)
            row = {}
            for o in OUTS:
                v = cell(calc, f"out_{o}").value
                if isinstance(v, str) and v in ERRORS:
                    row[o] = "ERROR:" + v
                else:
                    row[o] = None if v in ("–", None, "") else v
            results.append(row)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    json.dump(results, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
