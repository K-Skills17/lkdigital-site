"""Acceptance: clinic dashboard workbooks (spec section 6).

    python3 tests/dashboard_test.py   (after python3 dashboard/build_xlsx.py)
"""
import json
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent / "xlsx"))
from paths import DEMO_JSON, OUT  # noqa: E402
from xlsx_cache import ERRORS, recalculate_copy  # noqa: E402

DEMO_FILE = OUT / "dashboard-clinica-demo.xlsx"
BLANK_FILE = OUT / "dashboard-clinica.xlsx"
DEMO = json.loads(DEMO_JSON.read_text(encoding="utf-8"))
TABS = ["Como usar", "Entrada semanal", "Dashboard", "Por canal", "Metas", "Campanha A vs B"]
HEADERS = ["Semana (início)", "Canal", "Investimento", "Leads", "Qualificados", "Agendados",
           "Compareceram", "Fechados", "Receita fechada"]
INPUT_FILL = "F6EEDC"  # light gold

failures = 0


def check(ok, msg):
    global failures
    print(f"{'PASS' if ok else 'FAIL'}  {msg}")
    if not ok:
        failures += 1


def close(a, b, tol=0.006):
    return isinstance(a, (int, float)) and abs(a - b) <= tol


def recalculated(path):
    tmp = Path(tempfile.mkdtemp(prefix="dash-test-"))
    return load_workbook(recalculate_copy(path, tmp), data_only=True)


def error_cells(wb):
    return [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb.worksheets for row in ws.iter_rows()
            for c in row if isinstance(c.value, str) and c.value in ERRORS]


def fill(c):
    return (c.fill.fgColor.rgb or "")[-6:] if c.fill and c.fill.fill_type == "solid" else ""


for label, path in (("demo", DEMO_FILE), ("em branco", BLANK_FILE)):
    check(path.exists(), f"{label}: {path.name} exists")
    f = load_workbook(path)
    check(f.sheetnames == TABS, f"{label}: tabs are {', '.join(TABS)}")
    ent = f["Entrada semanal"]
    check([ent.cell(row=4, column=i + 1).value for i in range(9)] == HEADERS, f"{label}: weekly input columns match spec")

    dv = [d for d in ent.data_validations.dataValidation if "B5" in str(d.sqref)]
    check(bool(dv) and all(ch in dv[0].formula1 for ch in ["Meta", "Google", "GBP/orgânico", "Indicação", "Instagram orgânico", "Outro"]),
          f"{label}: Canal dropdown has the 6 channels")
    check(all(fill(ent.cell(row=r, column=c)) == INPUT_FILL for r in (5, 200, 404) for c in range(1, 10)),
          f"{label}: input cells have the light gold fill")
    check(ent.freeze_panes == "A5" and f["Dashboard"].freeze_panes == "A4", f"{label}: header rows frozen")
    check(404 - 5 + 1 >= 52 * 6, f"{label}: room for 52 weeks × 6 channels")
    check(all(str(ent.cell(row=r, column=10).value).startswith("=") and fill(ent.cell(row=r, column=10)) != INPUT_FILL for r in (5, 404)),
          f"{label}: non-input column is a formula without input fill")

    # Every calculated cell is a formula: no numbers typed into output areas.
    typed = []
    for sheet, area in (("Dashboard", "B3:I48"), ("Por canal", "B6:Q12"), ("Campanha A vs B", "B10:C12")):
        for row in f[sheet][area]:
            for c in row:
                if isinstance(c.value, (int, float)):
                    typed.append(f"{sheet}!{c.coordinate}")
    check(not typed, f"{label}: no hardcoded results in output areas{' ' + str(typed[:5]) if typed else ''}")

    texts = [str(c.value) for ws in f.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    check(not any(re.search(r"\bpre[çc]os?\b|a partir de|garantia", t, re.I) for t in texts),
          f"{label}: no \"preço\", \"a partir de\" or \"garantia\" in any cell")

    d = f["Dashboard"]
    check("R$" in d["B20"].number_format and "R$" in d["B6"].number_format, f"{label}: currency formatted R$")
    check(d["B3"].number_format.startswith("dd/mm/yyyy") and ent["A5"].number_format.startswith("dd/mm/yyyy"), f"{label}: dates dd/mm/aaaa")
    check(all('IFERROR(' in str(d[f"{c}{r}"].value) and '"–"' in str(d[f"{c}{r}"].value) for c in "BCD" for r in range(20, 27)),
          f"{label}: every derived ratio guarded with IFERROR → \"–\"")
    check(len(d.conditional_formatting) >= 1 and all("!" not in rule.formula[0] for cf in d.conditional_formatting for rule in cf.rules),
          f"{label}: conditional formatting present, same-sheet refs only (Google Sheets safe)")

    with zipfile.ZipFile(path) as z:
        charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
    check(len(charts) == 6, f"{label}: 6 charts (4 trend lines, funnel, channel mix) — found {len(charts)}")

    calc = recalculated(path)
    errs = error_cells(calc)
    check(not errs, f"{label}: zero formula errors after recalculation{' ' + str(errs[:5]) if errs else ''}")

    # Shipped cached values must equal a fresh recalculation.
    cached = load_workbook(path, data_only=True)
    diffs = []
    for ws in f.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    a, b = cached[ws.title][c.coordinate].value, calc[ws.title][c.coordinate].value
                    if a != b and not (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 1e-6) \
                            and not (a in (None, "") and b in (None, "")):
                        if hasattr(a, "date") and hasattr(b, "date") and a == b:
                            continue
                        diffs.append(f"{ws.title}!{c.coordinate}: {a!r} vs {b!r}")
    check(not diffs, f"{label}: cached values match recalculation{' ' + str(diffs[:3]) if diffs else ''}")

    if label == "demo":
        check(all(calc[t]["A1"].value == "Clínica fictícia — dados ilustrativos" for t in TABS),
              "demo: \"Clínica fictícia — dados ilustrativos\" on every tab")
        rows = [r for r in calc["Entrada semanal"].iter_rows(min_row=5, max_row=404, values_only=True) if r[0]]
        check(len(rows) == 40, f"demo: 8 weeks × 5 channels of data ({len(rows)} rows)")
        check(sum(r[2] for r in rows) == 6000 and sum(r[8] for r in rows) == 2 * 46750, "demo: 8-week sums = 2 × baseline month")
        weeks = sorted({r[0] for r in rows})
        windows_ok = True
        for w in range(len(weeks) - 3):
            win = [r for r in rows if r[0] in weeks[w:w + 4]]
            tot = [sum(r[k] for r in win) for k in (2, 3, 4, 5, 6, 7, 8)]
            windows_ok &= tot == [3000, 300, 120, 45, 30, 10, 46750]
        check(windows_ok, f"demo: every rolling 4-week window ({len(weeks) - 3}) equals the baseline month")

        dc = calc["Dashboard"]
        # Spec acceptance table (baseline month = "Mês até agora").
        exp = {"D20": ("CPL (só anúncios)", 10.00), "D22": ("Custo por agendamento", 66.67),
               "D23": ("Comparecimento", 0.6667), "D25": ("CAC (só anúncios)", 300), "D26": ("ROAS", 15.58)}
        for ref, (name, v) in exp.items():
            check(close(dc[ref].value, v, 0.01), f"demo acceptance: {name} = {v} (got {dc[ref].value!r})")
        check(close(dc["D21"].value, 25) and close(dc["D24"].value, 1 / 3), "demo: custo por qualificado 25, fechamento 33,3%")
        check(close(dc["B8"].value, 3000) and dc["C8"].value == 120 and dc["D8"].value == 45 and close(dc["E8"].value, 46750),
              "demo: top row month-to-date = 3.000 / 120 / 45 / 46.750")
        check(close(dc["B7"].value, 750) and close(dc["C7"].value, 30), "demo: top row 4-week average = 750 / 30")
        check(all(close(dc[f"C{r}"].value, dc[f"D{r}"].value) for r in range(20, 27)), "demo: last-4-weeks indicators = month-to-date (both baseline windows)")
        check(dc["B3"].value.date().isoformat() == "2026-09-28", "demo: last week = 28/09/2026")
        check([dc[f"B{r}"].value for r in range(44, 49)] == [300, 120, 45, 30, 10], "demo: funnel 300 → 120 → 45 → 30 → 10")
        trend = [dc[f"E{r}"].value for r in range(32, 40)]
        check(close(sum(trend), 2 * 46750), "demo: 8-week trend revenue sums to both months")

        pc = calc["Por canal"]
        total = {pc.cell(row=12, column=i).value for i in [1]}
        check(pc["A12"].value == "Total" and close(pc["H12"].value, dc["D17"].value) and pc["C12"].value == dc["D12"].value,
              "demo: Por canal totals match Dashboard month-to-date")
        by = {pc[f"A{r}"].value: (pc[f"C{r}"].value, pc[f"H{r}"].value) for r in range(6, 12)}
        check(max(by, key=lambda k: by[k][0]) == "Meta" and max(by, key=lambda k: by[k][1]) == "Indicação",
              "demo: Meta brings most leads, Indicação most revenue")
        check(pc["N11"].value == "–" and pc["O11"].value == "–", "demo: channel with no data shows \"–\"")

        ca = calc["Campanha A vs B"]
        check(ca["B11"].value == "–" and close(ca["C11"].value, 400) and ca["B12"].value == 0 and close(ca["C12"].value, 62.5),
              "demo: Campanha A 0 patients (\"–\"), B R$ 400 per patient, ROAS 62,5")
        check(close(ca["B10"].value, 2000) and close(ca["C10"].value, 1200), "demo: campaign spend 2.000 vs 1.200")
        m = calc["Dashboard"]
        check(m["E20"].value == 12 and m["F23"].value == "maior", "demo: Metas flow into the Dashboard")
    else:
        check(calc["Entrada semanal"]["A5"].value is None, "em branco: input table is empty")
        dc = calc["Dashboard"]
        check(all(dc[f"{c}{r}"].value == "–" for c in "BCD" for r in range(20, 27)), "em branco: derived metrics show \"–\"")
        check(all(dc[f"{c}{r}"].value in (0, None) for c in "BCDE" for r in (6, 7, 8)), "em branco: top row shows zeros")
        check(calc["Metas"]["B5"].value is None and calc["Campanha A vs B"]["B6"].value is None, "em branco: no demo targets or campaigns")
        check(all(calc[t]["A1"].value is None for t in TABS), "em branco: no fictional-data banner (no demo data)")

# Simulate a Monday: add a new week in the next empty rows of a copy; the dashboard must follow.
from datetime import datetime  # noqa: E402
wb = load_workbook(DEMO_FILE)
ent = wb["Entrada semanal"]
new = [("Meta", 500, 50, 20, 8, 6, 2, 10000), ("Google", 300, 20, 10, 4, 3, 1, 5000)]
for i, (canal, *vals) in enumerate(new):
    r = 45 + i
    ent.cell(row=r, column=1, value=datetime(2026, 10, 5))
    ent.cell(row=r, column=2, value=canal)
    for j, v in enumerate(vals):
        ent.cell(row=r, column=3 + j, value=v)
tmp = Path(tempfile.mkdtemp(prefix="dash-monday-")) / "monday.xlsx"
wb.save(tmp)
calc = recalculated(tmp)
dc = calc["Dashboard"]
check(dc["B3"].value.date().isoformat() == "2026-10-05", "new week: last week moves to 05/10/2026")
check(close(dc["B6"].value, 800) and dc["C6"].value == 30, "new week: top row last week = R$ 800 / 30 qualificados")
check(close(dc["B8"].value, 800) and close(dc["D25"].value, 800 / 3), "new week: month-to-date restarts in October (CAC 266,67)")
check(dc["G7"].value.date().isoformat() == "2026-09-14", "new week: 4-week window slides to 14/09")
check(not error_cells(calc), "new week: zero formula errors")

print(f"\n{failures} failing" if failures else "\nAll dashboard checks passed.")
sys.exit(1 if failures else 0)
