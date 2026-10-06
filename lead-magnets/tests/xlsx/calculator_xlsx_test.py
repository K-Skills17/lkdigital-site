"""Structure checks for public/ferramentas/arquivos/calculadora-cac.xlsx (numbers are covered by tests/calculator.test.cjs).

    python3 tests/calculator_xlsx_test.py
"""
import re
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent / "xlsx"))
from paths import DEMO_JSON, OUT  # noqa: E402
from xlsx_cache import ERRORS, recalculate_copy  # noqa: E402

PATH = OUT / "calculadora-cac.xlsx"
failures = 0


def check(ok, msg):
    global failures
    print(f"{'PASS' if ok else 'FAIL'}  {msg}")
    failures += 0 if ok else 1


f = load_workbook(PATH)
ws = f["Calculadora"]
check(f.sheetnames == ["Calculadora", "Como usar"], "tabs: Calculadora, Como usar")
check("Clínica fictícia — dados ilustrativos" in ws["A1"].value, "example labelled \"Clínica fictícia — dados ilustrativos\"")

outs = [n for n in f.defined_names if n.startswith("out_")]
not_formula = []
for n in outs:
    sheet, ref = next(iter(f.defined_names[n].destinations))
    v = f[sheet][ref.replace("$", "")].value
    if not (isinstance(v, str) and v.startswith("=")):
        not_formula.append(n)
check(len(outs) >= 25 and not not_formula, f"all {len(outs)} named outputs are formulas{' ' + str(not_formula) if not_formula else ''}")
ins = [n for n in f.defined_names if n.startswith("in_")]
fills = []
for n in ins:
    sheet, ref = next(iter(f.defined_names[n].destinations))
    c = f[sheet][ref.replace("$", "")]
    fills.append((c.fill.fgColor.rgb or "")[-6:] == "F6EEDC")
check(all(fills), f"all {len(ins)} named inputs have the input fill")
divs = [c.value for row in ws["M6:M31"] for c in row if isinstance(c.value, str) and "/" in c.value]
check(divs and all("IFERROR" in d for d in divs), "every division in the results is guarded with IFERROR")

texts = [str(c.value) for s in f.worksheets for row in s.iter_rows() for c in row if isinstance(c.value, str)]
check(not any(re.search(r"\bpre[çc]os?\b|a partir de|garantia", t, re.I) for t in texts), "no \"preço\", \"a partir de\" or \"garantia\"")

calc = load_workbook(recalculate_copy(PATH, Path(tempfile.mkdtemp())), data_only=True)
errs = [f"{s.title}!{c.coordinate}" for s in calc.worksheets for row in s.iter_rows() for c in row
        if isinstance(c.value, str) and c.value in ERRORS]
check(not errs, f"zero formula errors after recalculation{' ' + str(errs[:5]) if errs else ''}")
cached = load_workbook(PATH, data_only=True)
check(cached["Calculadora"]["M12"].value == 500, "cached values present (CAC real = 500 without recalculating)")

print(f"\n{failures} failing" if failures else "\nAll calculator xlsx checks passed.")
sys.exit(1 if failures else 0)
