"""Builds the clinic dashboard workbooks (spec section 6).

    python3 lead-magnets/xlsx/build_dashboard.py

Writes public/ferramentas/arquivos/dashboard-clinica-demo.xlsx (demo, 8 weeks of the fictional clinic)
and dashboard-clinica.xlsx (blank). Every calculated cell is a formula.
Cached values are filled in by xlsx_cache.py (LibreOffice recalculation) so
previews that don't calculate (phone quick look, WhatsApp) still show numbers.
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paths import DEMO_JSON, OUT, SITE  # noqa: E402
from xlsx_cache import fill_cached_values  # noqa: E402

DEMO = json.loads(DEMO_JSON.read_text(encoding="utf-8"))

CHANNELS = ["Meta", "Google", "GBP/orgânico", "Indicação", "Instagram orgânico", "Outro"]
FIRST, LAST = 5, 404  # input rows on "Entrada semanal"
IN = "'Entrada semanal'"

# Styles
FONT = "Arial"
F = lambda **k: Font(name=FONT, **k)  # noqa: E731
BLACK, RED, GREY = "0F0F0F", "7D6233", "6E6A63"  # ink, dark gold (text), muted
GOLD = "C5A368"
FILL_INPUT = PatternFill("solid", fgColor="F6EEDC")  # light gold
FILL_HEAD = PatternFill("solid", fgColor=BLACK)
FILL_BANNER = PatternFill("solid", fgColor="FFE680")
FILL_SOFT = PatternFill("solid", fgColor="F3F3F3")
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
FILL_YELLOW = PatternFill("solid", fgColor="FFEB9C")
FILL_RED = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

MONEY = '"R$" #,##0.00'
MONEY0 = '"R$" #,##0'
INT = "#,##0"
PCT = "0.0%"
RATIO = "0.0"
DATE = 'dd/mm/yyyy;;"–"'


def style_range(ws, ref, **kw):
    for row in ws[ref]:
        for c in row:
            for k, v in kw.items():
                setattr(c, k, v)


def head(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)
        c.border = BOX


def title(ws, text, demo, sub=None):
    if demo:
        ws["A1"] = "Clínica fictícia — dados ilustrativos"
        ws["A1"].font = F(bold=True, color="3A2E00")
        style_range(ws, "A1:I1", fill=FILL_BANNER)
    ws["A2"] = text
    ws["A2"].font = F(bold=True, size=16, color=BLACK)
    if sub:
        ws["A3"] = sub
        ws["A3"].font = F(italic=True, color=GREY)


def formula_cell(ws, ref, formula, fmt=None, **font):
    c = ws[ref]
    c.value = formula
    c.font = F(**font) if font else F()
    if fmt:
        c.number_format = fmt
    c.border = BOX
    return c


# ---------------------------------------------------------------- tabs

def tab_como_usar(wb, demo):
    ws = wb.active
    ws.title = "Como usar"
    title(ws, "Dashboard da clínica: como usar", demo, "10 minutos toda segunda-feira. Você só digita nas células douradas.")
    steps = [
        ("1", "Toda segunda-feira, abra a aba “Entrada semanal”."),
        ("2", "Adicione uma linha por canal com a data de início da semana (a segunda-feira) e escolha o canal na lista."),
        ("3", "Preencha só as células douradas: investimento em anúncios, leads, qualificados, agendados, compareceram, fechados e receita fechada."),
        ("4", "Abra o “Dashboard”: última semana, média de 4 semanas e mês até agora se atualizam sozinhos. Verde, amarelo e vermelho comparam com as suas “Metas”."),
        ("5", "Em “Por canal”, veja qual canal traz receita, não só leads. Anote uma decisão para a semana."),
    ]
    head(ws, 5, ["Passo", "O que fazer"])
    for i, (n, t) in enumerate(steps):
        r = 6 + i
        ws.cell(row=r, column=1, value=n).font = F(bold=True, size=12, color=RED)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
        c = ws.cell(row=r, column=2, value=t)
        c.font = F(size=11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 32

    ws["A12"] = "Legenda"
    ws["A12"].font = F(bold=True, size=12)
    ws["A13"].fill = FILL_INPUT
    ws["A13"].border = BOX
    ws["B13"] = "Célula dourada: você preenche."
    ws["A14"].border = BOX
    ws["B14"] = "Célula branca: fórmula. Não apague: é ela que calcula o dashboard."
    ws["A15"] = "–"
    ws["A15"].alignment = Alignment(horizontal="center")
    ws["B15"] = "Traço: não dá para calcular ainda (divisão por zero, por exemplo 0 fechados)."

    ws["A17"] = "Definições"
    ws["A17"].font = F(bold=True, size=12)
    defs = [
        ("Lead", "Qualquer contato novo: mensagem, ligação ou formulário."),
        ("Qualificado", "Lead com perfil e interesse real para a clínica (respondeu às perguntas de qualificação)."),
        ("Agendado", "Avaliação marcada na semana."),
        ("Compareceu", "Paciente que veio à avaliação."),
        ("Fechado", "Paciente que aprovou o plano de tratamento."),
        ("Receita fechada", "Soma dos tratamentos fechados na semana. Uso interno: nunca divulgue valores ao paciente."),
        ("Investimento", "Só anúncios (Meta, Google). Por isso CPL, CAC e ROAS aqui são “só anúncios”. Para o CAC real, com agência e equipe, use a Calculadora de CAC."),
    ]
    for i, (k, v) in enumerate(defs):
        r = 18 + i
        ws.cell(row=r, column=1, value=k).font = F(bold=True)
        c = ws.cell(row=r, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 28

    ws["A26"] = "Exemplo de linha"
    ws["A26"].font = F(bold=True, size=12)
    ex_head = ["Semana (início)", "Canal", "Investimento", "Leads", "Qualificados", "Agendados", "Compareceram", "Fechados", "Receita fechada"]
    head(ws, 27, ex_head)
    ex = [date(2026, 9, 7), "Meta", 450, 42, 14, 5, 3, 1, 350]
    fmts = [DATE, None, MONEY, INT, INT, INT, INT, INT, MONEY]
    for i, (v, f) in enumerate(zip(ex, fmts)):
        c = ws.cell(row=28, column=1 + i, value=v)
        c.font = F(color=GREY)
        c.border = BOX
        if f:
            c.number_format = f
    ws["A29"] = "Exemplo apenas para mostrar o formato; não entra nos cálculos."
    ws["A29"].font = F(italic=True, color=GREY, size=9)

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 95
    for col in "CDEFGHI":
        ws.column_dimensions[col].width = 13
    ws.sheet_view.showGridLines = False


def tab_entrada(wb, demo):
    ws = wb.create_sheet("Entrada semanal")
    title(ws, "Entrada semanal", demo, "Uma linha por semana por canal. Preencha só as células douradas.")
    cols = ["Semana (início)", "Canal", "Investimento", "Leads", "Qualificados", "Agendados", "Compareceram", "Fechados", "Receita fechada", "Verificação"]
    head(ws, 4, cols)
    ws.row_dimensions[4].height = 30
    widths = [15, 20, 14, 9, 13, 12, 14, 10, 16, 18]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    fmts = [DATE, None, MONEY, INT, INT, INT, INT, INT, MONEY]

    rows = DEMO["semanal"] if demo else []
    for r in range(FIRST, LAST + 1):
        for ci in range(9):
            c = ws.cell(row=r, column=ci + 1)
            c.fill = FILL_INPUT
            c.border = BOX
            c.font = F()
            if fmts[ci]:
                c.number_format = fmts[ci]
        # Funnel sanity check: warn, don't block.
        v = ws.cell(row=r, column=10, value=(
            f'=IF(A{r}="","",IF(AND(D{r}>=E{r},E{r}>=F{r},F{r}>=G{r},G{r}>=H{r}),"ok","confira o funil"))'))
        v.font = F(color=GREY)
        v.border = BOX
    for i, row in enumerate(rows):
        r = FIRST + i
        y, m, d = map(int, row["semana"].split("-"))
        vals = [date(y, m, d), row["canal"], row["investimento"], row["leads"], row["qualificados"],
                row["agendados"], row["compareceram"], row["fechados"], row["receita_fechada"]]
        for ci, v in enumerate(vals):
            ws.cell(row=r, column=ci + 1, value=v)

    dv_canal = DataValidation(type="list", formula1='"' + ",".join(CHANNELS) + '"', allow_blank=True,
                              showErrorMessage=True, errorTitle="Canal", error="Escolha um canal da lista.")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="43831", allow_blank=True,
                             showErrorMessage=True, errorTitle="Data", error="Use a data de início da semana (dd/mm/aaaa).")
    dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorTitle="Número", error="Use um número maior ou igual a zero.")
    dv_int = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorTitle="Número", error="Use um número inteiro maior ou igual a zero.")
    for dv in (dv_canal, dv_date, dv_num, dv_int):
        ws.add_data_validation(dv)
    dv_date.add(f"A{FIRST}:A{LAST}")
    dv_canal.add(f"B{FIRST}:B{LAST}")
    dv_num.add(f"C{FIRST}:C{LAST}")
    dv_num.add(f"I{FIRST}:I{LAST}")
    dv_int.add(f"D{FIRST}:H{LAST}")
    ws["J4"].comment = Comment("Avisa quando o funil não faz sentido (ex.: mais fechados que comparecimentos). Não bloqueia nada.", "LK Digital")
    ws.freeze_panes = f"A{FIRST}"

    names = {"e_semana": "A", "e_canal": "B", "e_invest": "C", "e_leads": "D", "e_qual": "E",
             "e_agend": "F", "e_comp": "G", "e_fech": "H", "e_receita": "I"}
    for name, col in names.items():
        wb.defined_names[name] = DefinedName(name, attr_text=f"{IN}!${col}${FIRST}:${col}${LAST}")


def tab_dashboard(wb, demo):
    ws = wb.create_sheet("Dashboard")
    title(ws, "Dashboard da clínica", demo)
    ws["A3"] = "Última semana lançada"
    ws["A3"].font = F(bold=True)
    formula_cell(ws, "B3", "=MAX(e_semana)", DATE, bold=True)
    ws["C3"] = "Primeira semana"
    ws["C3"].font = F(color=GREY)
    formula_cell(ws, "D3", "=MIN(e_semana)", DATE, color=GREY)

    # Top row: 4 numbers big x 3 periods. Period bounds in G:I.
    head(ws, 5, ["Período", "Investimento (anúncios)", "Leads qualificados", "Consultas agendadas", "Receita fechada"])
    head(ws, 5, ["De", "Até", "Semanas"], col=7)
    ws.row_dimensions[5].height = 30
    periods = [
        ("Última semana", "=$B$3", "=$B$3", "=IF($B$3=0,0,1)"),
        ("Média 4 semanas", "=IF($B$3=0,0,$B$3-21)", "=$B$3", "=IF($B$3=0,0,MIN(4,($B$3-$D$3)/7+1))"),
        ("Mês até agora", "=IF($B$3=0,0,DATE(YEAR($B$3),MONTH($B$3),1))", "=$B$3", "=IF($B$3=0,0,INT(($B$3-MAX($G$8,$D$3))/7)+1)"),
    ]
    for i, (label, de, ate, n) in enumerate(periods):
        r = 6 + i
        ws.cell(row=r, column=1, value=label).font = F(bold=True, size=11)
        ws.cell(row=r, column=1).border = BOX
        ws.cell(row=r, column=1).alignment = Alignment(vertical="center")
        formula_cell(ws, f"G{r}", de, DATE, color=GREY)
        formula_cell(ws, f"H{r}", ate, DATE, color=GREY)
        formula_cell(ws, f"I{r}", n, INT, color=GREY)
        ws.row_dimensions[r].height = 30
    # Big numbers read from the period totals below (rows 11, 13, 14, 17).
    big = [("B", 11, MONEY0), ("C", 13, INT), ("D", 14, INT), ("E", 17, MONEY0)]
    for col, src, fmt in big:
        formula_cell(ws, f"{col}6", f"=B{src}", fmt, bold=True, size=16)
        formula_cell(ws, f"{col}7", f"=IFERROR(C{src}/$I$7,0)", fmt, bold=True, size=16)
        formula_cell(ws, f"{col}8", f"=D{src}", fmt, bold=True, size=16)
        for r in (6, 7, 8):
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center", vertical="center")

    # Period totals.
    head(ws, 10, ["Totais do período", "Última semana", "Últimas 4 semanas", "Mês até agora"])
    totals = [("Investimento (anúncios)", "e_invest", MONEY), ("Leads", "e_leads", INT), ("Leads qualificados", "e_qual", INT),
              ("Consultas agendadas", "e_agend", INT), ("Comparecimentos", "e_comp", INT), ("Tratamentos fechados", "e_fech", INT),
              ("Receita fechada", "e_receita", MONEY)]
    for i, (label, rng, fmt) in enumerate(totals):
        r = 11 + i
        ws.cell(row=r, column=1, value=label).border = BOX
        ws.cell(row=r, column=1).font = F()
        for col, pr in (("B", 6), ("C", 7), ("D", 8)):
            formula_cell(ws, f"{col}{r}", f'=SUMIFS({rng},e_semana,">="&$G${pr},e_semana,"<="&$H${pr})', fmt)

    # Derived metrics, colored against Metas.
    head(ws, 19, ["Indicadores", "Última semana", "Últimas 4 semanas", "Mês até agora", "Meta", "Melhor se for"])
    derived = [
        ("CPL (só anúncios)", "{c}11/{c}12", MONEY),
        ("Custo por lead qualificado", "{c}11/{c}13", MONEY),
        ("Custo por consulta agendada", "{c}11/{c}14", MONEY),
        ("Taxa de comparecimento", "{c}15/{c}14", PCT),
        ("Taxa de fechamento", "{c}16/{c}15", PCT),
        ("CAC (só anúncios)", "{c}11/{c}16", MONEY),
        ("ROAS (receita ÷ investimento)", "{c}17/{c}11", RATIO),
    ]
    for i, (label, expr, fmt) in enumerate(derived):
        r = 20 + i
        ws.cell(row=r, column=1, value=label).border = BOX
        ws.cell(row=r, column=1).font = F(bold=True)
        for col in "BCD":
            formula_cell(ws, f"{col}{r}", f'=IFERROR({expr.format(c=col)},"–")', fmt, bold=True)
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center")
        formula_cell(ws, f"E{r}", f'=IF(Metas!B{5 + i}="","",Metas!B{5 + i})', fmt, color=GREY)
        formula_cell(ws, f"F{r}", f"=Metas!C{5 + i}", None, color=GREY)
        ws[f"E{r}"].alignment = Alignment(horizontal="center")
        ws[f"F{r}"].alignment = Alignment(horizontal="center")
    ws["A27"] = "Faixa amarela (tolerância)"
    ws["A27"].font = F(color=GREY)
    formula_cell(ws, "B27", "=Metas!B13", PCT, color=GREY)
    ws["A28"] = "Verde: dentro da meta · Amarelo: até a tolerância fora · Vermelho: além disso · Sem cor: sem meta definida."
    ws["A28"].font = F(italic=True, color=GREY, size=9)

    # Conditional formatting refers only to cells on this sheet (Google Sheets can't reference other sheets).
    rng = "B20:D26"
    ok = 'AND(ISNUMBER(B20),ISNUMBER($E20))'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({ok},IF($F20="menor",B20<=$E20,B20>=$E20))'], fill=FILL_GREEN, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({ok},IF($F20="menor",B20<=$E20*(1+$B$27),B20>=$E20*(1-$B$27)))'], fill=FILL_YELLOW, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[ok], fill=FILL_RED, stopIfTrue=True))

    # Weekly trend (last 8 weeks) and funnel data for the charts.
    ws["A30"] = "Tendência semanal (últimas 8 semanas)"
    ws["A30"].font = F(bold=True, size=12)
    head(ws, 31, ["Semana", "Investimento (anúncios)", "Leads qualificados", "Consultas agendadas", "Receita fechada"])
    for i in range(8):
        r = 32 + i
        back = 7 * (7 - i)
        formula_cell(ws, f"A{r}", f"=IF($B$3=0,0,$B$3-{back})" if back else "=$B$3", DATE)
        for col, rng_name, fmt in (("B", "e_invest", MONEY0), ("C", "e_qual", INT), ("D", "e_agend", INT), ("E", "e_receita", MONEY0)):
            formula_cell(ws, f"{col}{r}", f"=IF($A{r}=0,0,SUMIFS({rng_name},e_semana,$A{r}))", fmt)

    ws["A42"] = "Funil do mês até agora"
    ws["A42"].font = F(bold=True, size=12)
    head(ws, 43, ["Etapa", "Quantidade", "Conversão da etapa anterior"])
    funnel = [("Leads", 12), ("Qualificados", 13), ("Agendados", 14), ("Compareceram", 15), ("Fechados", 16)]
    for i, (label, src) in enumerate(funnel):
        r = 44 + i
        ws.cell(row=r, column=1, value=label).border = BOX
        formula_cell(ws, f"B{r}", f"=D{src}", INT)
        if i:
            formula_cell(ws, f"C{r}", f'=IFERROR(B{r}/B{r - 1},"–")', PCT)
        else:
            ws[f"C{r}"].border = BOX

    # Charts: one line chart per number (own scale each), plus the funnel.
    cats = Reference(ws, min_col=1, min_row=32, max_row=39)
    specs = [("B", "Investimento em anúncios por semana"), ("C", "Leads qualificados por semana"),
             ("D", "Consultas agendadas por semana"), ("E", "Receita fechada por semana")]
    for i, (col, name) in enumerate(specs):
        ch = LineChart()
        ch.title = name
        ch.style = 2
        ch.height, ch.width = 6.2, 11
        ch.legend = None
        cidx = ord(col) - 64
        ch.add_data(Reference(ws, min_col=cidx, min_row=31, max_row=39), titles_from_data=True)
        ch.set_categories(cats)
        s = ch.series[0]
        s.graphicalProperties.line.solidFill = BLACK
        s.graphicalProperties.line.width = 28000
        s.smooth = False
        ch.x_axis.number_format = "dd/mm"
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.majorGridlines = None
        anchor = ["K3", "Q3", "K16", "Q16"][i]
        ws.add_chart(ch, anchor)

    bar = BarChart()
    bar.type = "bar"
    bar.title = "Funil do mês"
    bar.style = 2
    bar.height, bar.width = 7.5, 11
    bar.legend = None
    bar.add_data(Reference(ws, min_col=2, min_row=43, max_row=48), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=44, max_row=48))
    bar.x_axis.scaling.orientation = "maxMin"  # Leads on top
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.y_axis.majorGridlines = None
    bar.series[0].graphicalProperties.solidFill = GOLD
    bar.series[0].graphicalProperties.line.solidFill = GOLD
    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = True
    bar.dataLabels.showSerName = False
    bar.dataLabels.showCatName = False
    bar.dataLabels.showLegendKey = False
    bar.dataLabels.showPercent = False
    ws.add_chart(bar, "K29")

    for col, w in zip("ABCDEFGHI", [30, 17, 17, 17, 17, 13, 12, 12, 9]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["J"].width = 3
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"


def tab_por_canal(wb, demo):
    ws = wb.create_sheet("Por canal")
    title(ws, "Por canal: quem traz receita, não só leads", demo)
    ws["A3"] = "Período"
    ws["A3"].font = F(bold=True)
    ws["B3"] = "Mês até agora"
    ws["B3"].fill = FILL_INPUT
    ws["B3"].border = BOX
    ws["B3"].font = F(bold=True)
    dv = DataValidation(type="list", formula1='"Mês até agora,Últimas 4 semanas,Última semana,Tudo"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    ws["C3"] = "de"
    ws["C3"].alignment = Alignment(horizontal="right")
    formula_cell(ws, "D3", '=IF($B$3="Tudo",Dashboard!$D$3,IF($B$3="Última semana",Dashboard!$G$6,IF($B$3="Últimas 4 semanas",Dashboard!$G$7,Dashboard!$G$8)))', DATE)
    ws["E3"] = "até"
    ws["E3"].alignment = Alignment(horizontal="right")
    formula_cell(ws, "F3", "=Dashboard!$B$3", DATE)

    cols = ["Canal", "Investimento", "Leads", "Qualificados", "Agendados", "Compareceram", "Fechados", "Receita fechada",
            "CPL", "Custo por qualificado", "Custo por agendamento", "Comparecimento", "Fechamento", "CAC", "ROAS",
            "% dos leads", "% da receita"]
    head(ws, 5, cols)
    ws.row_dimensions[5].height = 32
    sums = ["e_invest", "e_leads", "e_qual", "e_agend", "e_comp", "e_fech", "e_receita"]
    sum_fmt = [MONEY, INT, INT, INT, INT, INT, MONEY]
    total_row = 6 + len(CHANNELS)
    for i, ch in enumerate(CHANNELS + ["Total"]):
        r = 6 + i
        is_total = ch == "Total"
        c = ws.cell(row=r, column=1, value=ch)
        c.font = F(bold=True)
        c.border = BOX
        for j, (rng, fmt) in enumerate(zip(sums, sum_fmt)):
            col = chr(66 + j)
            if is_total:
                f = f"=SUM({col}6:{col}{total_row - 1})"
            else:
                f = f'=SUMIFS({rng},e_canal,$A{r},e_semana,">="&$D$3,e_semana,"<="&$F$3)'
            formula_cell(ws, f"{col}{r}", f, fmt, bold=is_total)
        ratios = [("I", f"B{r}/C{r}", MONEY), ("J", f"B{r}/D{r}", MONEY), ("K", f"B{r}/E{r}", MONEY),
                  ("L", f"F{r}/E{r}", PCT), ("M", f"G{r}/F{r}", PCT), ("N", f"B{r}/G{r}", MONEY),
                  ("O", f"H{r}/B{r}", RATIO), ("P", f"C{r}/C${total_row}", PCT), ("Q", f"H{r}/H${total_row}", PCT)]
        for col, expr, fmt in ratios:
            c = formula_cell(ws, f"{col}{r}", f'=IFERROR({expr},"–")', fmt, bold=is_total)
            c.alignment = Alignment(horizontal="right")
        if is_total:
            style_range(ws, f"A{r}:Q{r}", fill=FILL_SOFT)

    ws[f"A{total_row + 2}"] = ("Leitura: compare “% dos leads” com “% da receita”. Canal com muito lead e pouca receita "
                               "precisa de qualificação; canal com pouco lead e muita receita merece mais atenção.")
    ws[f"A{total_row + 2}"].font = F(italic=True, color=GREY)
    ws[f"A{total_row + 3}"] = "CPL, CAC e ROAS aqui consideram só o investimento em anúncios. Canais orgânicos aparecem com custo zero."
    ws[f"A{total_row + 3}"].font = F(italic=True, color=GREY)

    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "clustered"
    ch.title = "% dos leads x % da receita por canal"
    ch.style = 2
    ch.height, ch.width = 8.5, 16
    ch.add_data(Reference(ws, min_col=16, max_col=17, min_row=5, max_row=total_row - 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=6, max_row=total_row - 1))
    ch.x_axis.scaling.orientation = "maxMin"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.number_format = "0%"
    ch.y_axis.majorGridlines = None
    ch.series[0].graphicalProperties.solidFill = "9E9E9E"
    ch.series[0].graphicalProperties.line.solidFill = "9E9E9E"
    ch.series[1].graphicalProperties.solidFill = GOLD
    ch.series[1].graphicalProperties.line.solidFill = GOLD
    ch.legend.position = "b"
    ws.add_chart(ch, f"A{total_row + 5}")

    widths = [20, 14, 9, 12, 11, 13, 10, 15, 11, 13, 13, 14, 12, 11, 9, 11, 12]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B6"


def tab_metas(wb, demo):
    ws = wb.create_sheet("Metas")
    title(ws, "Metas", demo, "Defina as metas da sua clínica. O Dashboard fica verde, amarelo ou vermelho comparando com elas.")
    head(ws, 4, ["Indicador", "Meta", "Melhor se for", "Como pensar a meta"])
    rows = [
        ("CPL (só anúncios)", 12, "menor", MONEY, "Quanto você aceita pagar por um contato."),
        ("Custo por lead qualificado", 30, "menor", MONEY, "Custo por contato com perfil real."),
        ("Custo por consulta agendada", 55, "menor", MONEY, "Custo para colocar uma avaliação na agenda."),
        ("Taxa de comparecimento", 0.70, "maior", PCT, "Agendados que vieram. Lembretes melhoram isso."),
        ("Taxa de fechamento", 0.30, "maior", PCT, "Avaliações que viraram tratamento."),
        ("CAC (só anúncios)", 350, "menor", MONEY, "Custo de anúncios por paciente fechado."),
        ("ROAS (receita ÷ investimento)", 10, "maior", RATIO, "Receita fechada para cada real em anúncios."),
    ]
    dv = DataValidation(type="list", formula1='"menor,maior"', allow_blank=False)
    ws.add_data_validation(dv)
    for i, (label, meta, direction, fmt, hint) in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=1, value=label).font = F(bold=True)
        ws.cell(row=r, column=1).border = BOX
        m = ws.cell(row=r, column=2, value=meta if demo else None)
        m.number_format = fmt
        d = ws.cell(row=r, column=3, value=direction)
        for c in (m, d):
            c.fill = FILL_INPUT
            c.border = BOX
            c.font = F()
            c.alignment = Alignment(horizontal="center")
        dv.add(f"C{r}")
        ws.cell(row=r, column=4, value=hint).font = F(color=GREY)
    ws["A13"] = "Faixa amarela (tolerância)"
    ws["A13"].font = F(bold=True)
    ws["B13"] = 0.10
    ws["B13"].number_format = "0%"
    ws["B13"].fill = FILL_INPUT
    ws["B13"].border = BOX
    ws["B13"].alignment = Alignment(horizontal="center")
    ws["D13"] = "Ex.: 10% = amarelo quando está até 10% pior que a meta."
    ws["D13"].font = F(color=GREY)
    ws["A15"] = ("Use os números da sua própria clínica para definir metas. Comece pela média das últimas 4 semanas "
                 "e busque melhorar um pouco por mês.")
    ws["A15"].font = F(italic=True, color=GREY)
    for col, w in zip("ABCD", [32, 14, 14, 60]):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False


def tab_campanhas(wb, demo):
    ws = wb.create_sheet("Campanha A vs B")
    title(ws, "Campanha A vs B: clique barato ≠ paciente barato", demo,
          "Preencha as células douradas de duas campanhas e compare o custo por paciente fechado.")
    head(ws, 5, ["", "Campanha A", "Campanha B"])
    camp = DEMO["campanhas_ep5"]
    inputs = [("Cliques", "cliques", INT), ("CPC (custo por clique)", "cpc", MONEY),
              ("Pacientes fechados", "fechados", INT), ("Receita fechada", "receita_fechada", MONEY)]
    for i, (label, key, fmt) in enumerate(inputs):
        r = 6 + i
        ws.cell(row=r, column=1, value=label).font = F(bold=True)
        ws.cell(row=r, column=1).border = BOX
        for j, col in enumerate("BC"):
            c = ws[f"{col}{r}"]
            c.value = camp[j][key] if demo else None
            c.number_format = fmt
            c.fill = FILL_INPUT
            c.border = BOX
            c.font = F()
            c.alignment = Alignment(horizontal="center")
    outputs = [("Investimento (cliques × CPC)", "={c}6*{c}7", MONEY),
               ("Custo por paciente fechado", '=IFERROR({c}10/{c}8,"–")', MONEY),
               ("ROAS (receita ÷ investimento)", '=IFERROR({c}9/{c}10,"–")', RATIO)]
    for i, (label, f, fmt) in enumerate(outputs):
        r = 10 + i
        ws.cell(row=r, column=1, value=label).font = F(bold=True)
        ws.cell(row=r, column=1).border = BOX
        for col in "BC":
            c = formula_cell(ws, f"{col}{r}", f.format(c=col), fmt, bold=True, size=12)
            c.alignment = Alignment(horizontal="center")
    formula_cell(ws, "A14", '=IF(OR(B6="",C6=""),"Preencha as duas campanhas para comparar.",'
                 'IF(AND(ISNUMBER(B11),ISNUMBER(C11)),IF(B11<C11,"A Campanha A traz pacientes mais baratos.",'
                 'IF(C11<B11,"A Campanha B traz pacientes mais baratos.","Mesmo custo por paciente.")),'
                 'IF(ISNUMBER(C11),"A Campanha A ainda não fechou paciente: o clique barato não virou paciente.",'
                 'IF(ISNUMBER(B11),"A Campanha B ainda não fechou paciente: o clique barato não virou paciente.",'
                 '"Nenhuma das campanhas fechou paciente ainda."))))', None, bold=True, color=RED)
    ws.merge_cells("A14:C14")
    ws["A16"] = "Lição: CPC baixo não significa paciente barato. Compare custo por paciente fechado e ROAS, não o custo do clique."
    ws["A16"].font = F(italic=True, color=GREY)
    for col, w in zip("ABC", [34, 20, 20]):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False


def page_setup(wb, demo):
    """A4 landscape, one page wide; the Dashboard fits on one page."""
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1 if ws.title in ("Dashboard", "Metas", "Campanha A vs B", "Por canal") else 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.5
    wb["Dashboard"].print_area = "A1:W49"
    wb["Por canal"].print_area = "A1:Q34"
    rows = len(DEMO["semanal"]) if demo else 0
    wb["Entrada semanal"].print_area = f"A1:J{FIRST - 1 + max(rows, 40)}"
    wb["Entrada semanal"].print_title_rows = "4:4"


def build(demo: bool, out: Path):
    wb = Workbook()
    tab_como_usar(wb, demo)
    tab_entrada(wb, demo)
    tab_dashboard(wb, demo)
    tab_por_canal(wb, demo)
    tab_metas(wb, demo)
    tab_campanhas(wb, demo)
    page_setup(wb, demo)
    wb.active = wb.sheetnames.index("Dashboard") if demo else 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == wb.active.title
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    fill_cached_values(out)
    print(f"{out.relative_to(SITE)}")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build(True, OUT / "dashboard-clinica-demo.xlsx")
    build(False, OUT / "dashboard-clinica.xlsx")
