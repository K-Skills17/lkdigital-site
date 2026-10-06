"""Builds public/ferramentas/arquivos/calculadora-cac.xlsx: the CAC calculator as formulas (spec section 7).

    python3 calculator/build_xlsx.py

Same logic as calculator/calc.js; tests/calculator.test.cjs checks that both produce the
same numbers. Inputs and outputs are named cells (in_*, out_*) so tests can drive them.
Prefilled with the fictional demo clinic ("Clínica fictícia — dados ilustrativos").
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paths import DEMO_JSON, OUT, SITE  # noqa: E402
from xlsx_cache import fill_cached_values  # noqa: E402

DEMO = json.loads(DEMO_JSON.read_text(encoding="utf-8"))
SHEET = "Calculadora"
PROC_FIRST, PROC_LAST = 24, 33  # 10 procedure rows

FONT = "Arial"
F = lambda **k: Font(name=FONT, **k)  # noqa: E731
BLACK, RED, GREY = "0F0F0F", "7D6233", "6E6A63"  # ink, dark gold (text), muted
FILL_INPUT = PatternFill("solid", fgColor="F6EEDC")  # light gold
FILL_HEAD = PatternFill("solid", fgColor=BLACK)
FILL_BANNER = PatternFill("solid", fgColor="FFE680")
FILL_BEST = PatternFill("solid", fgColor="FFC7CE")
FILL_SOFT = PatternFill("solid", fgColor="F3F3F3")
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
MONEY = '"R$" #,##0.00'
MONEY0 = '"R$" #,##0'
INT = "#,##0"
PCT = "0.0%"
PCT0 = "0%"
RATIO = "0.0"


def name(wb, n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=f"'{SHEET}'!{ref}")


def section(ws, ref, text):
    c = ws[ref]
    c.value = text
    c.font = F(bold=True, color="FFFFFF")
    c.fill = FILL_HEAD
    row = c.row
    start = c.column
    for col in range(start, start + (2 if start == 1 else 7)):
        ws.cell(row=row, column=col).fill = FILL_HEAD


def label(ws, ref, text, hint=None, bold=False):
    c = ws[ref]
    c.value = text
    c.font = F(bold=bold)
    c.border = BOX
    c.alignment = Alignment(wrap_text=True, vertical="center")
    if hint:
        c.comment = Comment(hint, "LK Digital")


def inp(ws, ref, value, fmt=None):
    c = ws[ref]
    c.value = value
    c.fill = FILL_INPUT
    c.border = BOX
    c.font = F()
    if fmt:
        c.number_format = fmt
    return c


def out(ws, ref, formula, fmt=None, bold=True, size=10, color=None):
    c = ws[ref]
    c.value = formula
    c.font = F(bold=bold, size=size, color=color)
    c.border = BOX
    c.alignment = Alignment(horizontal="right", vertical="center")
    if fmt:
        c.number_format = fmt
    return c


def build(out_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    m = DEMO["mes_base"]
    ltv = DEMO["ltv_exemplo"]

    ws["A1"] = "Clínica fictícia — dados ilustrativos · Substitua pelos números da sua clínica (células douradas)."
    ws["A1"].font = F(bold=True, color="3A2E00")
    for col in range(1, 18):
        ws.cell(row=1, column=col).fill = FILL_BANNER
    ws["A2"] = "Calculadora de CAC da clínica"
    ws["A2"].font = F(bold=True, size=16)
    ws["A3"] = "Preencha só as células douradas. Tudo o que é branco é fórmula. “–” = não dá para calcular (divisão por zero)."
    ws["A3"].font = F(italic=True, color=GREY)

    # ---------------- inputs (A:B, procedures A:J) ----------------
    section(ws, "A5", "Custos do mês")
    rows = [
        (6, "Investimento em anúncios / mês", "in_invest", m["investimento_anuncios"], MONEY, None),
        (7, "Custos fixos de marketing / mês (agência, ferramentas, freelancers)", "in_fixos", m["custos_fixos_marketing"], MONEY, None),
        (8, "Salário da secretária (opcional)", "in_salario", None, MONEY, "Salário mensal com encargos."),
        (9, "% do tempo da secretária com leads", "in_pct_sec", None, PCT0, "Parte do tempo dela dedicada a responder e agendar leads."),
    ]
    for r, text, n, v, fmt, hint in rows:
        label(ws, f"A{r}", text, hint)
        inp(ws, f"B{r}", v, fmt)
        name(wb, n, f"$B${r}")
    label(ws, "A10", "Custo da secretária dedicado a leads")
    out(ws, "B10", "=B8*B9", MONEY, bold=False)
    label(ws, "A11", "Custo total de marketing", bold=True)
    out(ws, "B11", "=B6+B7+B10", MONEY)
    name(wb, "out_custo_total", "$B$11")

    section(ws, "A13", "Funil do mês")
    funnel = [(14, "Leads", "in_leads", m["leads"]), (15, "Qualificados", "in_qual", m["qualificados"]),
              (16, "Agendados", "in_agend", m["agendados"]), (17, "Compareceram", "in_comp", m["compareceram"]),
              (18, "Fechados (aprovaram o tratamento)", "in_fech", m["fechados"])]
    for r, text, n, v in funnel:
        label(ws, f"A{r}", text)
        inp(ws, f"B{r}", v, INT)
        name(wb, n, f"$B${r}")
    label(ws, "A19", "Verificação do funil")
    out(ws, "B19", '=IF(OR(B15>B14,B16>B15,B17>B16,B18>B17),"Confira: uma etapa tem mais gente que a anterior","ok")',
        bold=False, color=GREY)
    ws["B19"].alignment = Alignment(horizontal="left", wrap_text=True)

    ws["A21"] = "Procedimentos fechados no mês"
    ws["A21"].font = F(bold=True, size=11)
    ws["A22"] = "Ticket médio é o valor interno do tratamento, para calcular receita. Não divulgue valores ao paciente."
    ws["A22"].font = F(italic=True, color=GREY, size=9)
    heads = ["Procedimento", "Fechamentos", "Ticket médio", "Margem bruta", "Receita", "Lucro bruto",
             "Custo alocado", "CAC", "Lucro após CAC", "CAC ÷ ticket"]
    for i, h in enumerate(heads):
        c = ws.cell(row=23, column=1 + i, value=h)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.row_dimensions[23].height = 28
    ws["G23"].comment = Comment("Custo total de marketing dividido pela participação de cada procedimento nos fechamentos. "
                                "Ex.: 4 de 10 fechamentos recebem 40% do custo.", "LK Digital")
    procs = DEMO["procedimentos"]
    tot = PROC_LAST + 1
    for k, r in enumerate(range(PROC_FIRST, PROC_LAST + 1)):
        p = procs[k] if k < len(procs) else None
        inp(ws, f"A{r}", p["nome"] if p else None)
        inp(ws, f"B{r}", p["fechamentos"] if p else None, INT)
        inp(ws, f"C{r}", p["ticket_medio"] if p else None, MONEY0)
        inp(ws, f"D{r}", m["margem_bruta"] if p else None, PCT0)
        out(ws, f"E{r}", f'=IF(B{r}="","",B{r}*C{r})', MONEY0, bold=False)
        out(ws, f"F{r}", f'=IF(B{r}="","",B{r}*C{r}*D{r})', MONEY0, bold=False)
        out(ws, f"G{r}", f'=IF(B{r}="","",IFERROR($B$11*B{r}/$B${tot},"–"))', MONEY0, bold=False)
        out(ws, f"H{r}", f'=IF(B{r}="","",IFERROR(G{r}/B{r},"–"))', MONEY, bold=False)
        out(ws, f"I{r}", f'=IF(B{r}="","",IFERROR(F{r}-G{r},"–"))', MONEY0)
        out(ws, f"J{r}", f'=IF(B{r}="","",IFERROR(H{r}/C{r},"–"))', PCT0, bold=False)
    ws.conditional_formatting.add(f"I{PROC_FIRST}:I{PROC_LAST}",
                                  FormulaRule(formula=[f"AND(ISNUMBER(I{PROC_FIRST}),I{PROC_FIRST}<0)"], font=F(bold=True, color="B3261E")))
    label(ws, f"A{tot}", "Total", bold=True)
    for col, fmt in (("B", INT), ("E", MONEY0), ("F", MONEY0), ("G", MONEY0)):
        out(ws, f"{col}{tot}", f"=SUM({col}{PROC_FIRST}:{col}{PROC_LAST})", fmt)
    for col in "ABCDEFGHIJ":
        ws[f"{col}{tot}"].fill = FILL_SOFT
    label(ws, f"A{tot + 1}", "Verificação")
    out(ws, f"B{tot + 1}", f'=IF(AND(B{tot}>0,B18>0,B{tot}<>B18),"A soma dos fechamentos por procedimento é diferente dos fechados do funil","ok")',
        bold=False, color=GREY)
    ws[f"B{tot + 1}"].alignment = Alignment(horizontal="left")

    r0 = tot + 3  # LTV inputs
    section(ws, f"A{r0}", "Retorno e LTV (opcional)")
    ltv_rows = [(r0 + 1, "% de pacientes que voltam por ano", "in_volta", ltv["pct_retorno_ano"], PCT0),
                (r0 + 2, "Receita de retorno por paciente / ano", "in_retorno", ltv["receita_retorno_ano"], MONEY0),
                (r0 + 3, "Horizonte (1 a 3 anos)", "in_anos", ltv["horizonte_anos"], INT)]
    for r, text, n, v, fmt in ltv_rows:
        label(ws, f"A{r}", text)
        inp(ws, f"B{r}", v, fmt)
        name(wb, n, f"$B${r}")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=False,
                        showErrorMessage=True, error="Use 1, 2 ou 3 anos.")
    ws.add_data_validation(dv)
    dv.add(f"B{r0 + 3}")
    VOLTA, RETORNO, ANOS = f"B{r0 + 1}", f"B{r0 + 2}", f"B{r0 + 3}"

    num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                         showErrorMessage=True, error="Use um número maior ou igual a zero.")
    ws.add_data_validation(num)
    for ref in ("B6:B8", "B14:B18", f"B{PROC_FIRST}:C{PROC_LAST}", RETORNO):
        num.add(ref)
    pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                         showErrorMessage=True, error="Use uma porcentagem entre 0% e 100%.")
    ws.add_data_validation(pct)
    for ref in ("B9", f"D{PROC_FIRST}:D{PROC_LAST}", VOLTA):
        pct.add(ref)

    # ---------------- outputs (L:M, "Onde está o dinheiro" L:R) ----------------
    section(ws, "L5", "Resultados")
    outputs = [
        (6, "Custo por lead (CPL, só anúncios)", "out_cpl", '=IFERROR(B6/B14,"–")', MONEY),
        (7, "Custo por qualificado", "out_custo_qual", '=IFERROR(B6/B15,"–")', MONEY),
        (8, "Custo por agendamento", "out_custo_agend", '=IFERROR(B6/B16,"–")', MONEY),
        (9, "Custo por comparecimento", "out_custo_comp", '=IFERROR(B6/B17,"–")', MONEY),
        (11, "CAC só anúncios (anúncios ÷ fechados)", "out_cac_ads", '=IFERROR(B6/B18,"–")', MONEY),
        (12, "CAC real (todos os custos ÷ fechados)", "out_cac_real", '=IFERROR(B11/B18,"–")', MONEY),
        (13, "Diferença: quanto o painel de anúncios esconde", "out_cac_gap", '=IFERROR(M12-M11,"–")', MONEY),
        (15, "Lead → qualificado", "out_conv1", '=IFERROR(B15/B14,"–")', PCT),
        (16, "Qualificado → agendado", "out_conv2", '=IFERROR(B16/B15,"–")', PCT),
        (17, "Agendado → compareceu", "out_conv3", '=IFERROR(B17/B16,"–")', PCT),
        (18, "Compareceu → fechou", "out_conv4", '=IFERROR(B18/B17,"–")', PCT),
        (19, "Lead → fechamento", "out_conv_total", '=IFERROR(B18/B14,"–")', PCT),
        (20, "1 em cada X leads vira paciente", "out_um_em_cada", '=IFERROR(B14/B18,"–")', INT),
        (22, "Receita", "out_receita", f"=E{tot}", MONEY0),
        (23, "Lucro bruto", "out_lucro", f"=F{tot}", MONEY0),
        (24, "Custo total de marketing", None, "=B11", MONEY0),
        (25, "ROI sobre o custo total de marketing", "out_roi", '=IFERROR((M23-M24)/M24,"–")', PCT0),
        (26, "ROAS (receita ÷ anúncios)", "out_roas", '=IFERROR(M22/B6,"–")', RATIO),
        (28, "Ticket médio", "out_ticket_medio", f'=IFERROR(E{tot}/B{tot},"–")', MONEY),
        (29, "Margem média", "out_margem_media", f'=IFERROR(F{tot}/E{tot},"–")', PCT),
        (30, "LTV (lucro por paciente no horizonte)", "out_ltv", f'=IFERROR(M28*M29+{RETORNO}*M29*{VOLTA}*{ANOS},"–")', MONEY),
        (31, "LTV : CAC real", "out_ltv_cac", '=IFERROR(M30/M12,"–")', RATIO),
    ]
    for r, text, n, f, fmt in outputs:
        label(ws, f"L{r}", text)
        out(ws, f"M{r}", f, fmt, size=11 if r in (11, 12) else 10, color=RED if r == 12 else None)
        if n:
            name(wb, n, f"$M${r}")
    ws["L30"].comment = Comment("LTV = ticket médio × margem + receita de retorno por ano × margem × % que volta × anos.", "LK Digital")
    for r, t in ((10, ""), (14, "Conversão do funil"), (21, "Dinheiro"), (27, "LTV")):
        if t:
            ws[f"L{r}"] = t
            ws[f"L{r}"].font = F(bold=True, color=GREY)

    # Where the money is.
    ws["L33"] = "Onde está o dinheiro"
    ws["L33"].font = F(bold=True, size=12)
    label(ws, "L34", "Recuperar desta parte de quem se perde em cada etapa", "Premissa editável. Recuperar 20% de quem se perde em uma etapa, uma de cada vez.")
    inp(ws, "M34", 0.2, PCT0)
    name(wb, "in_recuperar", "$M$34")
    label(ws, "L35", "Aumento de anúncios para comparação")
    inp(ws, "M35", 0.2, PCT0)
    name(wb, "in_mais_anuncios", "$M$35")
    heads = ["Etapa", "Conversão atual", "Conversão melhorada", "Fechados", "CAC real", "Redução do CAC", "Posição"]
    for i, h in enumerate(heads):
        c = ws.cell(row=37, column=12 + i, value=h)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.row_dimensions[37].height = 28
    steps = ["Lead → qualificado", "Qualificado → agendado", "Agendado → compareceu", "Compareceu → fechou"]
    for k, text in enumerate(steps):
        r = 38 + k
        label(ws, f"L{r}", text, bold=True)
        out(ws, f"M{r}", f"=M{15 + k}", PCT, bold=False)
        out(ws, f"N{r}", f'=IFERROR(MIN(1,M{r}+$M$34*(1-M{r})),"–")', PCT, bold=False)
        out(ws, f"O{r}", f'=IFERROR(B18*N{r}/M{r},"–")', RATIO, bold=False)
        out(ws, f"P{r}", f'=IFERROR($B$11/O{r},"–")', MONEY)
        out(ws, f"Q{r}", f'=IFERROR($M$12-P{r},"–")', MONEY)
        out(ws, f"R{r}", f'=IFERROR(RANK(Q{r},$Q$38:$Q$41)+COUNTIF($Q$38:Q{r},Q{r})-1,"–")', INT, bold=False)
        ws[f"R{r}"].alignment = Alignment(horizontal="center")
        name(wb, f"out_onde{k + 1}", f"$P${r}")
    ws.conditional_formatting.add("L38:R41", FormulaRule(formula=["$R38=1"], fill=FILL_BEST))
    label(ws, "L42", "Comparação: mais anúncios, mesmo funil", bold=True)
    out(ws, "O42", '=B18*(1+$M$35)', RATIO, bold=False)
    out(ws, "P42", '=IFERROR((B6*(1+$M$35)+B7+B10)/O42,"–")', MONEY)
    out(ws, "Q42", '=IFERROR($M$12-P42,"–")', MONEY)
    name(wb, "out_mais_anuncios", "$P$42")
    for col in "MN":
        ws[f"{col}42"].border = BOX
    label(ws, "L44", "Maior alavanca", bold=True)
    out(ws, "M44", '=IFERROR(INDEX(L38:L41,MATCH(MAX(Q38:Q41),Q38:Q41,0)),"–")', None, color=RED, size=11)
    ws["M44"].alignment = Alignment(horizontal="left")
    ws.merge_cells("M44:P44")
    name(wb, "out_maior", "$M$44")
    ws["L45"] = "Arrume o funil antes de comprar mais anúncios. Só os seus números, sem médias de mercado."
    ws["L45"].font = F(italic=True, color=GREY)

    widths = {"A": 44, "B": 16, "C": 14, "D": 12, "E": 13, "F": 13, "G": 13, "H": 12, "I": 14, "J": 11, "K": 3,
              "L": 44, "M": 16, "N": 14, "O": 11, "P": 13, "Q": 14, "R": 9}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    # "Como usar" tab.
    how = wb.create_sheet("Como usar")
    how["A1"] = "Como usar a calculadora"
    how["A1"].font = F(bold=True, size=16)
    lines = [
        "1. Na aba “Calculadora”, troque os números nas células douradas pelos de um mês da sua clínica.",
        "2. Custos: anúncios, custos fixos de marketing (agência, ferramentas) e, se quiser, a parte do salário da secretária dedicada a leads.",
        "3. Funil: leads, qualificados, agendados, compareceram e fechados no mesmo mês.",
        "4. Procedimentos: quantos fechamentos de cada tipo, o ticket médio interno e a margem bruta.",
        "5. Leia os resultados à direita: CAC só anúncios × CAC real, custo por etapa, conversões, ROI, ROAS, LTV e “Onde está o dinheiro”.",
        "",
        "Como os números são calculados",
        "CAC só anúncios = investimento em anúncios ÷ fechados. CAC real = custo total de marketing ÷ fechados.",
        "CAC por procedimento: o custo total é dividido pela participação de cada procedimento nos fechamentos.",
        "ROI = (lucro bruto − custo total de marketing) ÷ custo total de marketing. ROAS = receita ÷ anúncios.",
        "LTV = ticket médio × margem + receita de retorno por ano × margem × % que volta × anos.",
        "Onde está o dinheiro: para cada etapa, uma de cada vez, recupera 20% das pessoas que se perdem nela e recalcula o CAC real.",
        "(Aumentar todas as etapas em 20% relativos daria o mesmo CAC para qualquer etapa, porque os fechados são o produto das conversões.)",
        "A linha “mais anúncios” mostra o CAC real se você aumentar os anúncios em 20% com o mesmo funil.",
        "",
        "“–” significa que não dá para calcular ainda (por exemplo, zero fechados).",
        "Os números de exemplo são de uma clínica fictícia — dados ilustrativos.",
    ]
    for i, t in enumerate(lines):
        c = how.cell(row=3 + i, column=1, value=t)
        c.font = F(bold=t in ("Como os números são calculados",))
        c.alignment = Alignment(wrap_text=True, vertical="top")
    how.column_dimensions["A"].width = 120
    how.sheet_view.showGridLines = False

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    fill_cached_values(out_path)
    print(out_path.relative_to(SITE))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build(OUT / "calculadora-cac.xlsx")
