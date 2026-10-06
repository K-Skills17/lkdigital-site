"""Fills cached formula values into an openpyxl-written workbook.

openpyxl writes formulas without results, so viewers that don't calculate (phone
previews, WhatsApp, some email clients) show empty cells. This recalculates a copy with
LibreOffice, then writes each result into the original file's <v> element. The original
workbook (charts, validation, conditional formatting) is otherwise untouched, and Excel
and Google Sheets still recalculate on open.

    python3 scripts/xlsx_cache.py file.xlsx
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook

ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}


def recalculate_copy(src: Path, workdir: Path) -> Path:
    """Returns a LibreOffice-recalculated copy of src."""
    profile = workdir / "profile"
    out = workdir / "out"
    out.mkdir()
    cmd = ["soffice", f"-env:UserInstallation={profile.as_uri()}", "--headless", "--calc",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)]
    env = {**os.environ, "SAL_USE_VCLPLUGIN": "svp"}  # headless rendering backend
    subprocess.run(cmd, check=True, capture_output=True, timeout=180, env=env)
    result = out / src.name
    if not result.exists():
        raise RuntimeError(f"LibreOffice did not produce {result}")
    return result


def sheet_paths(zf: zipfile.ZipFile) -> dict:
    """Maps sheet name -> xml path inside the package."""
    wb = zf.read("xl/workbook.xml").decode()
    rels = zf.read("xl/_rels/workbook.xml.rels").decode()
    targets = dict(re.findall(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    targets.update({k: v for v, k in re.findall(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    out = {}
    for name, rid in re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
        t = targets[rid].lstrip("/")
        out[name.replace("&amp;", "&")] = t if t.startswith("xl/") else "xl/" + t
    return out


CELL_RE = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f>(?:<v>.*?</v>|<v\s*/>)?</c>', re.S)


def cached_xml(value):
    """Returns (type attribute, <v> text) for a computed value."""
    if isinstance(value, bool):
        return ' t="b"', "1" if value else "0"
    if isinstance(value, (int, float)):
        return "", repr(float(value)) if isinstance(value, float) else str(value)
    if hasattr(value, "toordinal"):  # date/datetime from a date-formatted formula
        from openpyxl.utils.datetime import to_excel
        return "", repr(float(to_excel(value)))
    text = "" if value is None else str(value)
    if text in ERRORS:
        return ' t="e"', text
    return ' t="str"', escape(text)


def fill_cached_values(path: Path) -> dict:
    """Recalculates path and writes cached values in place. Returns {errors: [...], formulas: n}."""
    path = Path(path)
    with tempfile.TemporaryDirectory(prefix="xlsx-cache-") as tmp:
        tmp = Path(tmp)
        calc = recalculate_copy(path, tmp)
        values = load_workbook(calc, data_only=True)
        errors, count = [], 0

        with zipfile.ZipFile(path) as zin:
            paths = sheet_paths(zin)
            by_path = {p: n for n, p in paths.items()}
            staged = tmp / "staged.xlsx"
            with zipfile.ZipFile(staged, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename in by_path:
                        ws = values[by_path[item.filename]]

                        def repl(m):
                            nonlocal count
                            ref, attrs, formula = m.groups()
                            v = ws[ref].value
                            count += 1
                            if isinstance(v, str) and v in ERRORS:
                                errors.append(f"{ws.title}!{ref} {v}")
                            attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
                            t, text = cached_xml(v)
                            return f'<c r="{ref}"{attrs}{t}><f>{formula}</f><v>{text}</v></c>'

                        data = CELL_RE.sub(repl, data.decode()).encode()
                    zout.writestr(item, data)
        shutil.move(staged, path)
    return {"formulas": count, "errors": errors}


if __name__ == "__main__":
    for p in sys.argv[1:]:
        r = fill_cached_values(Path(p))
        print(p, r["formulas"], "formulas,", len(r["errors"]), "errors", r["errors"][:10])
